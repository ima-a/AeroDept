"""Build a 13-slide editable PPTX and a companion PDF layout preview.

Run: PYTHONPATH=/tmp/aerodepth_pptx_deps python presentations/aerodepth_lite/build_presentation.py
The PDF is rendered from this shared layout, not exported by PowerPoint.
"""
from pathlib import Path
import json, csv, math, textwrap, zipfile
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'assets'
NAVY='102638'; INK='193447'; TEAL='008C83'; MINT='DFF3EE'
ORANGE='B35B15'; SAND='FFF0DF'; PAPER='F5F8FA'; WHITE='FFFFFF'
MUTED='526575'; LINE='DCE5EB'; LIGHT='EAF0F4'
FONT='/usr/share/fonts/noto/NotoSans-Regular.ttf'
BOLD='/usr/share/fonts/noto/NotoSans-Bold.ttf'
prs=Presentation(); prs.slide_width=Inches(16); prs.slide_height=Inches(9)
prs.core_properties.title='AeroDepth-Lite | Project Presentation'
prs.core_properties.subject='Monocular UAV range prediction and causal inference scheduling'
prs.core_properties.author='AeroDepth-Lite Project Team'
prs.core_properties.keywords='MidAir, MobileNetV3, depth regression, scheduling, validation'
previews=[]; overflow=[]; slides_text=[]

def rgb(c): return RGBColor.from_string(c)
def font(size,bold=False,symbols=False):
    path=('/usr/share/fonts/liberation/LiberationSans-Bold.ttf' if bold else '/usr/share/fonts/liberation/LiberationSans-Regular.ttf') if symbols else (BOLD if bold else FONT)
    return ImageFont.truetype(path,round(size*100/72))
def wrap(s,w,size,bold=False):
    f=font(size,bold,any(c in str(s) for c in '→↓↑')); out=[]
    for para in str(s).split('\n'):
        line=''
        for word in para.split():
            candidate=(line+' '+word).strip()
            if f.getlength(candidate)>w*100-3 and line: out.append(line);line=word
            else: line=candidate
        out.append(line)
    return out

class Page:
    def __init__(self,n,title,subtitle='',source='',dark=False):
        self.n=n; self.dark=dark; self.s=prs.slides.add_slide(prs.slide_layouts[6]); self.items=[]
        self.im=Image.new('RGB',(1600,900),'#'+(NAVY if dark else PAPER)); self.d=ImageDraw.Draw(self.im)
        self.s.background.fill.solid(); self.s.background.fill.fore_color.rgb=rgb(NAVY if dark else PAPER)
        if not dark:
            self.text(.65,.32,13,.25,'AERODEPTH-LITE  /  DEEP LEARNING PROJECT',10,TEAL,True)
            self.text(.65,.85,14.7,.68,title,30,INK,True)
            if subtitle:self.text(.65,1.58,14.6,.6,subtitle,15,MUTED)
            self.rect(.65,8.28,14.7,.013,LINE)
            self.text(.65,8.43,13.8,.31,source,9,MUTED)
            self.text(14.5,8.38,.85,.35,f'{n:02d} / 13',10,MUTED,align='right')
    def rect(self,x,y,w,h,fill,line=None,radius=False):
        sh=self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        if radius:sh.adjustments[0]=.08
        sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
        if line:sh.line.color.rgb=rgb(line)
        else:sh.line.fill.background()
        self.d.rectangle((x*100,y*100,(x+w)*100,(y+h)*100),fill='#'+fill,outline='#'+line if line else None)
        return sh
    def text(self,x,y,w,h,s,size=18,color=INK,bold=False,align='left',url=None):
        symbols=any(c in str(s) for c in '→↓↑')
        lines=wrap(s,w,size,bold); content='\n'.join(lines); self.items.append(content)
        sh=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        step=size*1.19/72
        if len(lines)*step>h+.055:overflow.append((self.n, s[:65],round(len(lines)*step,2),h))
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
            p.font.name='Arial' if symbols else 'Noto Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
            p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(size*1.19)
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            if url:
                for r in p.runs:r.hyperlink.address=url
            f=font(size,bold,symbols);tx=x*100
            if align=='center':tx+=(w*100-f.getlength(line))/2
            if align=='right':tx+=w*100-f.getlength(line)
            self.d.text((tx,(y+i*step)*100),line,font=f,fill='#'+color,anchor='lt')
        return sh
    def arrow(self,x1,y1,x2,y2,color=TEAL):
        sh=self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        sh.line.color.rgb=rgb(color);sh.line.width=Pt(2)
        e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','med');e.set('len','med');sh.line._get_or_add_ln().append(e)
        self.d.line((x1*100,y1*100,x2*100,y2*100),fill='#'+color,width=3)
        a=math.atan2(y2-y1,x2-x1);points=[(x2*100,y2*100)]
        for da in [-.48,.48]:points.append((x2*100-13*math.cos(a+da),y2*100-13*math.sin(a+da)))
        self.d.polygon(points,fill='#'+color)
    def picture(self,path,x,y,w,h):
        self.s.shapes.add_picture(str(path),Inches(x),Inches(y),Inches(w),Inches(h))
        im=Image.open(path).convert('RGB').resize((round(w*100),round(h*100)),Image.Resampling.LANCZOS)
        self.im.paste(im,(round(x*100),round(y*100)))
    def band(self,y,title,body,kind='teal'):
        self.rect(.65,y,14.7,.88,MINT if kind=='teal' else SAND)
        self.text(.9,y+.12,14.2,.3,title,16,TEAL if kind=='teal' else ORANGE,True)
        self.text(.9,y+.49,14.2,.3,body,13,INK)
    def card(self,x,y,w,h,kicker,title,body,accent=TEAL):
        self.rect(x,y,w,h,WHITE);self.rect(x,y,.055,h,accent)
        self.text(x+.25,y+.2,w-.5,.28,kicker,11,accent,True)
        self.text(x+.25,y+.68,w-.5,.85,title,22,INK,True)
        self.text(x+.25,y+1.65,w-.5,h-1.8,body,17,MUTED)
    def table(self,x,y,width,cols,headers,rows,rowh=.7,size=16):
        for ri,row in enumerate([headers]+rows):
            xx=x
            for ci,val in enumerate(row):
                ww=width*cols[ci];fill=NAVY if ri==0 else (WHITE if ri%2 else LIGHT)
                self.rect(xx,y+ri*rowh,ww,rowh,fill)
                self.text(xx+.14,y+ri*rowh+.14,ww-.28,rowh-.2,str(val),size,WHITE if ri==0 else INK,ri==0)
                xx+=ww
    def finish(self,notes):
        self.s.notes_slide.notes_text_frame.text=notes
        self.im.save(ASSETS/f'slide_{self.n:02d}.png');previews.append(self.im)
        slides_text.append({'slide':self.n,'text':self.items,'speaker_notes':notes})

# 1 — Title, with team fields left explicitly editable.
p=Page(1,'',dark=True)
p.rect(.7,.65,.7,.08,'3CD6BE');p.text(.7,.96,11,.4,'DEEP LEARNING PROJECT PRESENTATION',13,'76DAC9',True)
p.text(.7,1.75,10.9,1.1,'AeroDepth-Lite',49,WHITE,True)
p.text(.75,3.04,10.4,1.35,'Monocular Drone-Camera Range Prediction\nand Adaptive Inference Scheduling',25,'DCEAF0')
p.text(.75,4.72,9.8,.65,'A compact depth model. A causal update policy.\nAn evidence-based study of accuracy and compute.',17,'A6BFCE')
p.rect(.75,6.0,10.0,2.27,'1A3549')
p.text(1,6.2,9.4,.4,'TEAM MEMBERS  •  EDIT THESE FIELDS',11,'76DAC9',True)
p.text(1,6.7,9.3,.5,'[Member name — Roll no.]   |   [Member name — Roll no.]',16,WHITE)
p.text(1,7.3,9.3,.7,'[Department]  •  [College / University]\nGuide: [Name and designation]',15,'DCEAF0')
for i,(a,b,c) in enumerate([('RGB','Single frame','76DAC9'),('RANGE','Range map · metres','FFFFFF'),('UPDATE','Infer or reuse','F3B979')]):
    yy=1.9+i*1.8;p.rect(11.6,yy,3.55,1.25,'1A3549');p.text(11.85,yy+.18,3.05,.35,a,20,c,True);p.text(11.85,yy+.68,3.05,.3,b,13,'DCEAF0')
    if i<2:p.arrow(13.38,yy+1.29,13.38,yy+1.69,'76DAC9')
p.text(11.6,7.7,3.5,.4,'PROJECT REVIEW  /  01',11,'A6BFCE')
p.finish('Introduce the team after replacing the title-slide placeholders. AeroDepth-Lite predicts dense Euclidean camera range in metres and evaluates causal inference scheduling. This deck follows the supplied 13-slide outline. It summarizes existing saved research; no experiment was rerun. Source: docs/FINAL_REPORT.md. Evidence current to 29 October 2026.')

# 2 — Problem.
p=Page(2,'Problem statement and motivation','How can a small drone estimate scene range while limiting repeated inference?','Source: docs/FINAL_REPORT.md · README.md')
p.card(.65,2.5,4.7,3.65,'01  /  PERCEPTION','One image,\nuncertain range','Recover dense metric range from monocular RGB under changing terrain, weather and visibility.')
p.card(5.65,2.5,4.7,3.65,'02  /  COMPUTE','Repeated inference\nhas a cost','A fresh model call for every frame adds work. Reusing an old map can increase staleness and miss scene changes.')
p.card(10.65,2.5,4.7,3.65,'03  /  EVALUATION','Average error can\nhide near errors','Overall gains may coexist with worse errors below 10 m. Model calls, update delay and measured time all matter.',ORANGE)
p.band(6.7,'Project gap','Evaluate compact range prediction and causal scheduling together, including near-field regressions and real timing.')
p.finish('The gap is the integration and evaluation question addressed by this project, not a claim that no previous work studied efficient depth. Practical motivation is limited compute on small UAV platforms. This project uses simulation and does not demonstrate flight safety, obstacle detection or battery savings. Source: docs/FINAL_REPORT.md, Problem and objectives; Limitations and conclusion.')

# 3 — Objectives.
p=Page(3,'Objectives','Measure prediction quality, update behaviour and actual execution cost.','Source: docs/FINAL_REPORT.md · results/presentation_validation_20261029T155521Z/REPORT.md')
objectives=[('01','Predict dense range','Use a compact MobileNetV3-Small encoder and skip decoder to produce a 192 × 320 range map.'),('02','Compare learning strategies','Evaluate the original student, expanded supervised A and exploratory knowledge distillation (KD).'),('03','Schedule causally','Use current RGB, simulated IMU, held GPS and prior predictions to choose fresh inference or reuse.'),('04','Verify the trade-offs','Report overall / near error, model calls, update delay, measured timing and limitations.')]
for i,(num,title,body) in enumerate(objectives):
    y=2.42+i*1.13;p.rect(.65,y,.7,.7,TEAL);p.text(.65,y+.16,.7,.4,num,19,WHITE,True,'center');p.text(1.65,y,4.1,.5,title,20,INK,True);p.text(5.8,y,9.3,.85,body,18,MUTED)
p.band(7.12,'Expected outputs','Reproducible model exports, frozen split manifests, comparative evidence and a local demonstration dashboard.')
p.finish('Objectives are not guarantees of improvement. The saved deliverables include original epoch-4 and supervised-A epoch-9 inference exports, manifest identities, comparison tables and a local Streamlit dashboard. The reserved independent final evaluation remains incomplete. Sources: README.md; docs/FINAL_REPORT.md.')

# 4 — Literature table.
p=Page(4,'Literature review and research gap','Method foundations and the limitations relevant to this project.','References: Eigen (2014), Howard (2019), He (2016), Hinton (2015), Fonder & Van Droogenbroeck (2019). Links in notes.')
p.table(.65,2.35,14.7,[.23,.34,.43],['Study','Model / approach','Limit relative to our project scope'],[
    ['Eigen et al., 2014','Coarse-to-fine monocular depth','NYU / KITTI results do not establish UAV scheduler behaviour.'],
    ['Howard et al., 2019','MobileNetV3; hardware-aware design','Efficient encoder; metric range needs a decoder and task-specific training.'],
    ['He et al., 2016','Residual learning / ResNet','Classification backbone; larger capacity does not guarantee a suitable depth teacher.'],
    ['Hinton et al., 2015','Teacher-to-student distillation','Knowledge transfer does not guarantee improved near-range regression.'],
    ['Fonder & Van Droogenbroeck, 2019','MidAir synthetic UAV dataset','Rich simulation data; real-flight generalization must be tested separately.']
],rowh=.72,size=15)
p.band(7.0,'Research gap addressed','Study the joint range-quality / update-cost trade-off on a fixed UAV subset; retain negative findings and per-path checks.')
p.finish('References verified against primary sources. Eigen, Puhrsch & Fergus: Depth Map Prediction from a Single Image using a Multi-Scale Deep Network (2014), https://arxiv.org/abs/1406.2283 . Howard et al.: Searching for MobileNetV3 (2019), https://arxiv.org/abs/1905.02244 . He et al.: Deep Residual Learning for Image Recognition (CVPR 2016; preprint 2015), https://arxiv.org/abs/1512.03385 . Hinton, Vinyals & Dean: Distilling the Knowledge in a Neural Network (2015 preprint), https://arxiv.org/abs/1503.02531 . Fonder & Van Droogenbroeck: Mid-Air: A multi-modal dataset for extremely low altitude drone flights (CVPRW 2019), https://orbi.uliege.be/handle/2268/234665 . The right-hand column is our assessment of applicability, not verbatim claims by the authors. These studies motivate components; their reported benchmark scores are not directly compared with this custom subset. No field-wide novelty or state-of-the-art claim is made.')

# 5 — Contributions.
p=Page(5,'Novelty and team contributions','The contribution is the implemented combination and its controlled evaluation.','Source: src/aerodepth_lite/ · docs/FINAL_REPORT.md · saved student and scheduler comparisons')
p.card(.65,2.4,7.1,2.45,'METHOD','Compact metric-range model','1,027,489 parameters; MobileNetV3-Small with a custom skip decoder.',TEAL)
p.card(8.05,2.4,7.3,2.45,'SYSTEM','Causal multi-signal scheduler','Combines visual change, simulated motion, prior predicted risk and map age.',TEAL)
p.card(.65,5.1,7.1,2.45,'EXPERIMENT','Matched A / KD comparison','Same starting student state and training budget; a failed teacher gate stays visible.',ORANGE)
p.card(8.05,5.1,7.3,2.45,'REPRODUCIBILITY','Auditable delivery','Grouped splits, frozen artifacts, near-field metrics, ablations and measured timing.',TEAL)
p.finish('Project-specific contributions: implement a compact positive metric-range decoder, integrate a causal scheduler, perform controlled A/KD experiments, preserve rejected-teacher and negative-KD evidence, and publish reproducible inference/evaluation artifacts with a local dashboard. MobileNetV3, ResNet, MidAir and distillation are prior work. The scheduler risk cache is verified to preserve decisions and reduces scheduler overhead, but does not show a robust end-to-end gain. Individual team responsibilities were not supplied, so no personal contribution is invented. Source: docs/FINAL_REPORT.md.')

# 6 — Dataset.
p=Page(6,'Dataset details','MidAir · synthetic multi-modal data for low-altitude drone flights','Sources: expanded_preparation_20261026/{manifest_train_val.json.gz,data_validation.json} · MidAir paper and FAQ')
p.table(.65,2.35,9.1,[.34,.2,.19,.27],['Project split','Frames','Paths','Status'],[
 ['Expanded train','44,098','12','Completed'],['Validation','4,410','2','Selection / analysis'],['Reserved new test','8,820*','4','NOT RUN']
],rowh=.66,size=16)
p.text(.65,5.2,9.05,.5,'48,508 available train + validation pairs',22,TEAL,True)
p.text(.65,5.88,9.0,1.5,'Train: Kite 17,640 · PLE 26,458\nValidation: Kite 2,205 · PLE 2,205\nClasses / per-class counts: not applicable (dense regression).',16,INK)
p.rect(10.05,2.35,5.3,5.25,WHITE)
p.text(10.3,2.6,4.8,.4,'FORMAT & DATA QUALITY',13,TEAL,True)
p.text(10.3,3.17,4.8,3.9,'RGB: 1024 × 1024 JPEG\nRange: uint16 PNG / float16 bits\nSensors: HDF5 (IMU / GPS)\n\n0 duplicate RGB paths in manifest.\nInvalid range pixels are masked.\nNear-field evaluation is imbalanced: 97.23% of near pixels are from PLE.',16,INK)
p.text(.65,7.35,9.0,.65,'*Expected from metadata; 8 RGB/depth archives missing. Legacy 4,436-frame test is historical and excluded from this split.',12,MUTED)
p.text(10.3,7.07,4.8,.35,'Dataset source / download',12,TEAL,True,url='https://midair.ulg.ac.be/download.html')
p.finish('Official full dataset: more than 420k video frames, 54 trajectories, https://orbi.uliege.be/handle/2268/234665 . Our custom subset is NOT the official benchmark. There are 20 training climate realizations across 12 paths and 2 validation realizations across 2 paths. Exact counts: train 44,098, validation 4,410, total available 48,508. New reserved expected 8,820 frames across 4 paths, not acquired/evaluated. Exact split counts are used instead of a random percentage split. Source: results/expanded_preparation_20261026/manifest_train_val.json.gz and reservations.json. Duplicate RGB path count was recomputed from that frozen manifest: zero; this is not a perceptual duplicate-image audit. Pixel masking handles out-of-range values; integrity checks cover filenames, pair matching and bounded decode checks, not a claim of zero invalid pixels. Near-pixel imbalance from docs/FINAL_REPORT.md. Dataset range is Euclidean distance to focal point: https://midair.ulg.ac.be/faq.html . Download: https://midair.ulg.ac.be/download.html . Attribution: Fonder & Van Droogenbroeck, CVPRW 2019; see NOTICE.md for dataset terms.')

# 7 — Preprocessing.
p=Page(7,'Data preprocessing and leakage controls','Preserve the range target and keep related flight paths in one split.','Source: docs/expanded_training.md · expanded_preparation_20261026/dataset_contract.json · src/aerodepth_lite/data.py')
steps=[('01','Validate pairs','Match RGB, range and sensor records using the frozen allowlist.'),('02','Decode & mask','Reinterpret PNG bits as float16 range. Keep finite GT in [0.1, 1000] m.'),('03','Resize & normalize','Full-frame 192 × 320; bilinear RGB; nearest target / mask; ImageNet normalization once.'),('04','Augment training','Paired horizontal flip for RGB, target and mask; no validation augmentation.')]
for i,(n,t,b) in enumerate(steps):
 y=2.4+i*1.2;p.text(.72,y,.6,.4,n,19,TEAL,True);p.text(1.5,y,4.25,.5,t,20,INK,True);p.text(5.9,y,9.0,.9,b,17,MUTED)
p.band(7.18,'Leakage controls','Group weather variants of a path; quarantine legacy tests; select on validation only; keep current GT out of scheduler decisions.')
p.finish('Invalid depth values are replaced by zero only after creating the valid-pixel mask, then excluded from losses and metrics. Near field is valid ground-truth range <10 m. No crop or learned scale correction; softplus+0.1 has no upper prediction clamp. Feature extraction is learned by the encoder; there is no separate handcrafted feature selection for the depth network. Weather-equivalent paths identified using position-stream hashes stay grouped. Different paths may share scenery, so geographic isolation is not established. Validation was used repeatedly for development and model selection. Current GT range and privileged true velocity do not enter V2 decisions. Sources: docs/FINAL_REPORT.md; docs/expanded_training.md; results/expanded_preparation_20261026/dataset_contract.json.')

# 8 — Editable diagram.
p=Page(8,'Proposed system architecture','Causal runtime workflow · drawn with editable shapes and connectors','Source: src/aerodepth_lite/{model.py,scheduler_v2.py,presentation_scheduler.py} · docs/FINAL_REPORT.md')
def box(x,y,w,h,title,body,fill=WHITE):
 p.rect(x,y,w,h,fill,LINE);p.text(x+.18,y+.18,w-.36,.45,title,18,TEAL,True,'center');p.text(x+.18,y+.76,w-.36,h-.86,body,14,INK,False,'center')
box(.65,2.55,2.25,1.65,'Current RGB','Frame t')
box(3.35,2.55,2.6,1.65,'Preprocess','Resize + normalize')
box(6.45,2.55,2.65,1.65,'V2 decision','Visual + motion\nrisk + age',MINT)
box(9.65,2.2,2.65,1.55,'INFER','Student model',MINT)
box(9.65,4.1,2.65,1.55,'REUSE','Last range map',SAND)
box(12.85,3.1,2.5,1.8,'Range output','Fresh / reused\nsource age')
p.arrow(2.92,3.37,3.28,3.37);p.arrow(5.97,3.37,6.38,3.37)
p.arrow(9.12,3.1,9.58,2.97);p.arrow(9.12,3.7,9.58,4.85)
p.arrow(12.32,2.97,12.78,3.65);p.arrow(12.32,4.85,12.78,4.4)
box(.65,5.2,5.3,1.6,'Simulated navigation','IMU + held GPS velocity')
p.arrow(5.97,6.0,7.35,4.28)
box(6.45,6.1,5.85,1.45,'State / cache','Last inferred RGB + prediction + age')
p.arrow(7.8,6.03,7.8,4.28)
p.arrow(14.1,4.97,14.1,6.82);p.arrow(14.1,6.82,12.38,6.82)
p.text(.65,7.45,5.3,.55,'First frame: infer. Maximum reuse age: 4 frames.',13,MUTED)
p.text(12.85,7.08,2.5,.72,'GT range is used for offline scoring only.',12,ORANGE,True)
p.finish('Diagram consists of ordinary editable PowerPoint shapes and arrow connectors; no image-generation model was used. Runtime: resize current RGB; scheduler computes change versus last inferred RGB and combines causal simulated IMU/GPS, last predicted risk and age; infer updates cached prediction, otherwise reuse the prior map without warping. The model normalizes RGB internally exactly once (the preprocessing box denotes the full transform contract). State resets at path boundaries. Risk cache is refreshed only on fresh-source token/reset. Ground truth is used offline, never as a decision input. Source: docs/FINAL_REPORT.md and implementation. Teacher/KD is training-only, not on this inference path.')

# 9 — Model and methodology.
p=Page(9,'Deep learning model and methodology','Small student for inference; larger teacher evaluated only as a training aid.','Source: docs/FINAL_REPORT.md · frozen launch configurations · src/aerodepth_lite/model.py')
p.rect(.65,2.3,14.7,1.15,NAVY)
p.text(.95,2.51,14.1,.64,'RGB → MobileNetV3-Small → 576→128 projection → skip decoder → 1-channel range',19,WHITE,True,'center')
p.text(.65,3.7,7.0,.4,'STUDENT ARCHITECTURE',13,TEAL,True)
p.text(.65,4.25,7.0,1.5,'1,027,489 parameters\nSkip indices 0 / 1 / 3 / 8\nDecoder widths 64 / 32 / 24 / 16\nBilinear upsampling; softplus + 0.1 m',18,INK)
p.text(.65,6.12,7.0,.94,'Why selected: compact ImageNet-initialized encoder with a custom dense decoder. Teacher: ResNet-50 V1.5.',17,MUTED)
p.text(.65,7.06,7.0,.3,'Tools: Python · PyTorch / torchvision · NumPy · Streamlit',12,TEAL)
p.text(8.05,3.7,7.3,.4,'TRAINING CONFIGURATION',13,TEAL,True)
p.text(8.05,4.25,7.25,2.7,'A / KD: 10 epochs each; selected epoch 9\nTeacher B: 20 epochs; selected epoch 3\nAdamW · batch 2 · seed 42 · clip 1.0\nStudent LR 10⁻⁴ · weight decay 10⁻⁴\nCosine minimum: 10⁻⁶\nTeacher encoder / decoder LR: 10⁻⁵ / 10⁻⁴',17,INK)
p.rect(.65,7.47,14.7,.56,MINT)
p.text(.85,7.59,14.2,.33,'Loss = L1 range + 0.2 L1 log-range + 0.1 gradient L1  |  KD adds 0.2 log-range consistency',14,TEAL,True)
p.finish('Tools: Python, PyTorch, torchvision, NumPy, Pillow, h5py, Matplotlib and Streamlit. Hardware recorded: GTX 1650 (4 GB) for training; CPU/CUDA timing was measured separately. Original student completed 50 epochs on 17,635 training frames and selected epoch 4 by recorded AMP validation AbsRel. A and exploratory KD independently start from identical original epoch-4 weights and BN buffers with fresh matched optimizer/scheduler/scaler/RNG state. Their common dataset is expanded 44,098 training frames; final comparisons use FP32 validation. A and KD selected epoch 9 by minimum pooled FP32 validation AbsRel. Teacher is ImageNet1K_V2 initialized ResNet-50 V1.5, selected epoch 3; it failed the near-field acceptance gate. During exploratory KD it is frozen, eval mode, FP32, no-grad, and sees the same augmented image as the student. Near loss weight is zero. KD adds 0.2 masked absolute log-range consistency on valid training pixels. No classification-logit temperature KD is used. Sources: docs/FINAL_REPORT.md; configs and frozen launch copies. Numerical AMP recovery preserves update accounting and retries rather than silently skipping steps.')

# 10 — Training and evaluation curves.
p=Page(10,'Experimental results: training and validation','Recorded curves from the completed expanded supervised and exploratory KD runs.','Source: results/presentation_validation_20261029T155521Z/training_curves.json · plots regenerated from saved values')
p.picture(ASSETS/'training_curves.png',.65,2.25,14.7,4.26)
p.rect(.65,6.76,7.1,1.18,MINT);p.text(.9,6.92,6.6,.33,'Both students selected epoch 9',18,TEAL,True);p.text(.9,7.4,6.6,.35,'Selection metric: pooled FP32 validation AbsRel.',14,INK)
p.rect(8.05,6.76,7.3,1.18,SAND);p.text(8.3,6.92,6.8,.33,'Lower training loss is not proof of generalization',16,ORANGE,True);p.text(8.3,7.4,6.8,.35,'A and KD objectives differ; compare validation error.',14,INK)
p.finish('These are actual saved per-epoch values, not synthetic curves and not a rerun. Both 10-epoch runs start independently from original epoch 4. Left: train objective; KD contains an extra term, so loss magnitude is not directly comparable across objectives. Right: same FP32 validation AbsRel, permitting comparison on the familiar development set. Original 50-epoch training overfits after its selected epoch 4; its recorded curve uses AMP validation and should not be merged as an identical FP32 series. Source: results/presentation_validation_20261029T155521Z/training_curves.json. Selected checkpoints were evaluated using matched masks and valid-pixel counts.')

# 11 — Model comparison.
rows=list(csv.DictReader((ROOT/'results/student_validation_comparison_20261029T152648Z/comparison.csv').open()))
pooled={r['model']:r for r in rows if r['scope']=='pooled'}
p=Page(11,'Experimental results: model comparison','4,410 validation frames · lower error is better · identical evaluation masks','Source: results/student_validation_comparison_20261029T152648Z/comparison.csv · docs/FINAL_REPORT.md')
data=[]
for key,label in [('epoch4','Original · epoch 4'),('student_A','Supervised A · epoch 9'),('exploratory_KD','Exploratory KD · epoch 9')]:
 r=pooled[key];data.append([label,f"{float(r['absrel']):.6f}",f"{float(r['rmse']):.3f}",f"{100*float(r['delta1']):.2f}%",f"{float(r['near_rmse']):.3f}"])
p.table(.65,2.32,14.7,[.30,.18,.18,.16,.18],['Model','AbsRel ↓','RMSE (m) ↓','δ1 ↑','Near RMSE ↓'],data,rowh=.76,size=17)
p.text(.65,5.4,14.7,.3,'Teacher B failed its acceptance gate: Kite near RMSE 6.636 m vs A 4.695 m.',13,ORANGE,True)
for x,w,val,lab,fill,col in [(.65,4.7,'11.84%','Lower AbsRel · A vs original',MINT,TEAL),(5.65,4.7,'16.18%','Higher near RMSE · A vs original',SAND,ORANGE),(10.65,4.7,'1.59%','Higher AbsRel · KD vs A',SAND,ORANGE)]:
 p.rect(x,5.75,w,1.26,fill);p.text(x+.23,5.86,w-.46,.6,val,29,col,True);p.text(x+.23,6.57,w-.46,.3,lab,12,INK)
p.text(.65,7.35,14.7,.65,'δ1 = fraction with max(prediction / GT, GT / prediction) < 1.25. Classification precision, recall, F1 and confusion matrix do not apply.',14,MUTED)
p.finish('AbsRel = mean(|P−GT|/GT); RMSE = sqrt(sum((P−GT)^2)/valid-pixel count). Values are pooled by additive pixel sums, not averaged per-frame RMSE. Valid GT is finite in [0.1,1000] m; near GT <10 m. All three models have 212,903,814 valid and 32,267,840 near pixels. A improves overall RMSE by 5.80% and AbsRel by 11.84% vs original, but near RMSE worsens 16.18%. Original and A differ in training data and budget, so this does not isolate an architecture effect. KD vs A is a matched one-seed comparison; KD pooled AbsRel is 1.59% worse and near RMSE 1.42% worse. Teacher B failed the gate: Kite foggy near RMSE 6.636090 m vs A 4.695176 m. The exploratory rejected-teacher run does not show that all distillation fails. No confidence interval, multi-seed statistical claim, independent new-test result or cross-paper benchmark superiority is available. Source: comparison.csv and docs/FINAL_REPORT.md.')

# 12 — Scheduling and timing.
p=Page(12,'Experimental results: scheduling and measured cost','Supervised A on validation · policy parameters frozen before this comparison','Source: presentation_validation_20261029T155521Z/{scheduler_comparison.csv,timing_individual.csv} · docs/FINAL_REPORT.md')
p.table(.65,2.3,9.15,[.30,.19,.28,.23],['Policy','Calls','Near RMSE (m)','Max delay¹'],[
 ['Every frame','4,410','3.906','0'],['Fixed budget','2,456','4.161','1'],['V2 adaptive','2,456','3.780','2'],['V2: no visual','1,789','5.529','4']
],rowh=.64,size=15)
p.rect(10.1,2.3,5.25,3.2,MINT)
p.text(10.35,2.54,4.75,.65,'9.14%',33,TEAL,True)
p.text(10.35,3.32,4.75,1.85,'Lower pooled near RMSE than fixed budget at equal total calls.\n\nOverall error slightly worsens; max event-update delay rises.',17,INK)
p.text(.65,5.84,14.7,.5,'Actual timing · 256-frame stress windows · median end-to-end seconds',19,INK,True)
p.table(.65,6.48,9.15,[.30,.23,.23,.24],['Device','Every frame','Fixed budget','V2 cached'],[['CPU','10.420','8.673','10.429'],['GTX 1650','9.018','7.863','8.765']],rowh=.48,size=13)
p.text(10.1,5.85,5.25,1.95,'Risk caching reduced scheduler time by 13.43% on CPU and 15.81% on GPU.\nNo robust end-to-end speedup was measured.',16,ORANGE,True)
p.finish('¹ Delay is in camera frames and refers to time until a prediction update after an offline-defined event, NOT obstacle-detection accuracy. Near RMSE V2 vs fixed budget: 9.14% lower pooled, but Kite near RMSE is slightly worse and PLE improves. V2 uses 905/1551 calls across Kite/PLE; budget fixed uses 1228/1228. Total calls coincide, not per-path calls. V2 overall AbsRel 0.268088 vs fixed 0.267810; max delay 2 vs 1. Visual ablation changes budget and state trajectory, so it is not an equal-compute causal estimate. Timing uses two fixed 128-frame windows, 3 repetitions/device, FP32 batch 1, CPU 2 threads, 10 excluded warmups. In these windows every-frame/fixed-budget/V2 each make 256/144/222 calls per run, unlike the full-path validation call budgets. CPU median [min,max] seconds: every 10.4203 [10.2725,14.3289], budget 8.6725 [8.3859,9.1672], cached V2 10.4290 [10.1152,10.6138]. CUDA: every 9.0176 [8.8663,9.4765], budget 7.8625 [7.7724,8.3073], cached V2 8.7650 [8.7047,9.3071]. Risk caching vs uncached improves scheduler median 13.43% CPU /15.81% GPU; end-to-end CPU is 3.07% worse and GPU 0.37% better with overlapping repetition ranges. Model-only latency is not complete throughput. No energy or battery benefit was measured.')

# 13 — Limits and conclusion.
p=Page(13,'Limitations, future scope and conclusion','Overall validation gains are useful only when the near-field and compute trade-offs stay visible.','Source: docs/FINAL_REPORT.md · results/acquisition_publication_20261029T180214Z/REPORT.md')
p.card(.65,2.4,7.1,3.45,'CURRENT LIMITATIONS','Evidence remains bounded','Two familiar validation paths; correlated frames; synthetic scenes; one seed; substantial depth error; near-field regression.',ORANGE)
p.card(8.05,2.4,7.3,3.45,'FUTURE SCOPE','Complete independent evaluation','Acquire the 8 missing reserved streams, freeze the test protocol, evaluate new paths, then study near-field losses and real-flight transfer.',TEAL)
p.rect(.65,6.2,14.7,1.72,NAVY)
p.text(.95,6.42,14.1,.38,'CONCLUSION',13,'76DAC9',True)
p.text(.95,6.94,14.0,.75,'Supervised A improves overall range prediction, while near-field error worsens.\nExploratory KD underperforms A; adaptive scheduling offers a mixed quality / cost trade-off.',20,WHITE,True)
p.finish('Reserved independent final study: NOT RUN because 8 image/depth streams are unavailable; sensor metadata alone is insufficient. The older 4,436-frame test belongs to the historical original-model study and is quarantined, not evidence for new-model generalization. Future work is proposed, not completed: acquisition and integrity checks, a frozen independent protocol, multi-seed replication, near-field loss/coverage experiments, real drone sensor latency and power measurements. Teacher-regression mechanisms are hypotheses; the separate gradient probe did not justify repair. No measured energy saving, battery gain, flight safety or guaranteed obstacle detection is claimed. Finish with supported findings and invite questions. Sources: docs/FINAL_REPORT.md and acquisition inventory.')

assert len(prs.slides)==13
path=OUT/'AeroDepth_Lite_Project_Presentation.pptx';prs.save(path)
previews[0].save(OUT/'AeroDepth_Lite_Preview.pdf',save_all=True,append_images=previews[1:],resolution=100.0)
thumb=Image.new('RGB',(1600,math.ceil(13/3)*318),'#DCE5EB')
for i,im in enumerate(previews):thumb.paste(im.resize((512,288)),(16+(i%3)*528,16+(i//3)*318))
thumb.save(OUT/'slide_overview.jpg',quality=90)
(OUT/'slide_content_and_notes.json').write_text(json.dumps(slides_text,indent=2,ensure_ascii=False))
with zipfile.ZipFile(path) as z:
    bad=z.testzip(); assert bad is None
    from lxml import etree
    for name in z.namelist():
        if name.endswith('.xml') or name.endswith('.rels'):etree.fromstring(z.read(name))
check=Presentation(path);assert len(check.slides)==13
(OUT/'layout_check.json').write_text(json.dumps({'slides':len(check.slides),'text_overflows':overflow,'xml_validation':'passed','preview_note':'PDF and thumbnails use the shared layout renderer, not PowerPoint export.'},indent=2))
print(json.dumps({'pptx':str(path),'slide_count':len(prs.slides),'text_overflows':overflow},indent=2))
